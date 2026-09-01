from io import BytesIO
from urllib.parse import parse_qs, urlparse

from PIL import Image
from docx.document import Document as DocxDocument
from docx.image.exceptions import UnrecognizedImageError
from pypdf import PdfReader

from app.models.course import ArticleImageAsset, ArticleText, AudioAsset, Course, CourseStatus, SourceType
from app.models.file_resource import FileResource, ResourceKind, ResourceStatus, ResourceVariant
from app.models.generation_job import GenerationJob, JobStatus
from app.services.course_export_service import CourseExportService
from app.services.job_service import build_course_export_fingerprint


def create_exportable_course(db_session):
    course = Course(
        user_id="test_user",
        title="可下载课程",
        source_type=SourceType.URL_IMPORT,
        status=CourseStatus.NEEDS_REVIEW,
    )
    db_session.add(course)
    db_session.flush()
    article_text = ArticleText(
        course_id=course.id,
        version=1,
        text="可下载课程\n\n第一句。第二句。",
        content_markdown="# 可下载课程\n\n第一句。第二句。",
    )
    db_session.add(article_text)
    db_session.commit()
    return course, article_text


def test_course_markdown_export_returns_attachment(client, db_session):
    course, _ = create_exportable_course(db_session)

    response = client.get(f"/courses/{course.id}/exports/markdown")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert "attachment" in response.headers["content-disposition"]
    assert ".md" in response.headers["content-disposition"]
    assert "# 可下载课程" in response.text
    assert "第一句。第二句。" in response.text


def test_course_word_and_pdf_exports_return_downloadable_files(client, db_session):
    course, _ = create_exportable_course(db_session)

    docx_response = client.get(f"/courses/{course.id}/exports/docx")
    pdf_response = client.get(f"/courses/{course.id}/exports/pdf")

    assert docx_response.status_code == 200
    assert docx_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert docx_response.content[:2] == b"PK"
    assert ".docx" in docx_response.headers["content-disposition"]

    assert pdf_response.status_code == 200
    assert pdf_response.headers["content-type"].startswith("application/pdf")
    assert pdf_response.content.startswith(b"%PDF")
    assert ".pdf" in pdf_response.headers["content-disposition"]


def test_course_pdf_export_preserves_chinese_code_block_text(client, db_session):
    course, article_text = create_exportable_course(db_session)
    article_text.content_markdown = """# 可下载课程

```text
下面是我准备做用户访谈时想问的问题：
[贴上你的问题]
请你用 The Mom Test 的原则帮我改写：
1. 删掉意见型问题
2. 删掉假设未来的问题
3. 尽量改成围绕过去真实行为、已有替代方案和已付成本的提问
4. 最后整理成一套 8-10 个可以直接访谈的问题清单
```
"""
    db_session.commit()

    response = client.get(f"/courses/{course.id}/exports/pdf")

    assert response.status_code == 200
    extracted_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages
    )
    assert "下面是我准备做用户访谈时想问的问题" in extracted_text
    assert "8-10 个可以直接访谈的问题清单" in extracted_text


def test_course_download_request_queues_export_job(client, db_session, monkeypatch):
    course, article_text = create_exportable_course(db_session)
    queued = {}
    monkeypatch.setattr(
        "app.services.job_service.enqueue_generation_job",
        lambda job_id: queued.update({"job_id": job_id}) or "queued-job-id",
    )

    response = client.post(f"/courses/{course.id}/downloads/markdown")

    assert response.status_code == 202
    body = response.json()
    assert body["status"] == "pending"
    assert body["job_type"] == "course_export_markdown"
    assert body["job_id"] == queued["job_id"]
    job = db_session.query(GenerationJob).filter(GenerationJob.id == body["job_id"]).one()
    assert job.course_id == course.id
    assert job.target_type == "course"
    assert job.target_id == course.id


def test_course_download_request_returns_ready_link_for_existing_resource(client, db_session):
    course, article_text = create_exportable_course(db_session)
    fingerprint = build_course_export_fingerprint(course, "markdown", article_text)
    resource = FileResource(
        user_id=course.user_id,
        owner_type="course",
        owner_id=course.id,
        resource_kind=ResourceKind.EXPORT,
        resource_variant=ResourceVariant.MARKDOWN,
        status=ResourceStatus.READY,
        source_fingerprint=fingerprint,
        title=course.title,
        filename="ready.md",
        storage_backend="r2",
        bucket="pagealong-media",
        object_key="downloads/course.md",
        object_path="https://media.pagealong.test/downloads/course.md",
        content_type="text/markdown; charset=utf-8",
        byte_size=12,
        checksum_sha256="a" * 64,
        metadata_json="{}",
    )
    db_session.add(resource)
    db_session.commit()

    response = client.post(f"/courses/{course.id}/downloads/markdown")

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ready"
    parsed = urlparse(body["download_url"])
    assert f"{parsed.scheme}://{parsed.netloc}{parsed.path}" == resource.object_path
    query = parse_qs(parsed.query)
    assert query["response-content-disposition"][0].startswith("attachment;")
    assert query["response-content-type"][0] == resource.content_type
    assert body["resource_id"] == resource.id
    assert body["job_id"] is None


def test_course_audio_download_request_returns_ready_link_when_resource_exists(client, db_session):
    course, article_text = create_exportable_course(db_session)
    resource = FileResource(
        user_id=course.user_id,
        owner_type="course",
        owner_id=course.id,
        resource_kind=ResourceKind.AUDIO,
        resource_variant=ResourceVariant.AUDIO,
        status=ResourceStatus.READY,
        source_fingerprint="audio-fingerprint",
        title=course.title,
        filename="ready.mp3",
        storage_backend="r2",
        bucket="pagealong-media",
        object_key="audio/course_1.mp3",
        object_path="https://media.pagealong.test/audio/course_1.mp3",
        content_type="audio/mpeg",
        byte_size=12,
        checksum_sha256="a" * 64,
        metadata_json="{}",
    )
    db_session.add(resource)
    db_session.flush()
    course.status = CourseStatus.READY
    course.current_audio_resource_id = resource.id
    db_session.commit()

    response = client.post(f"/courses/{course.id}/downloads/audio")

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ready"
    parsed = urlparse(body["download_url"])
    assert f"{parsed.scheme}://{parsed.netloc}{parsed.path}" == resource.object_path
    query = parse_qs(parsed.query)
    assert query["response-content-disposition"][0].startswith("attachment;")
    assert query["response-content-type"][0] == resource.content_type
    assert body["resource_id"] == resource.id
    assert body["job_id"] is None


def test_course_audio_download_request_marks_generation_job_failed_when_enqueue_fails(
    client, db_session, monkeypatch
):
    def fail_enqueue(course_id, job_id):
        raise RuntimeError("redis unavailable")

    monkeypatch.setattr("app.services.course_service.enqueue_audio_generation", fail_enqueue)
    course, _ = create_exportable_course(db_session)
    course.status = CourseStatus.TEXT_READY
    db_session.commit()

    response = client.post(f"/courses/{course.id}/downloads/audio")

    assert response.status_code == 503
    db_session.expire_all()
    course = db_session.get(Course, course.id)
    job = db_session.query(GenerationJob).filter(GenerationJob.course_id == course.id).one()
    assert course.status == CourseStatus.FAILED
    assert job.status == JobStatus.FAILED
    assert job.error_code == "queue_unavailable"


def test_course_docx_export_converts_unrecognized_images(db_session, tmp_path, monkeypatch):
    course, article_text = create_exportable_course(db_session)
    image_path = tmp_path / "cover.webp"
    image_buffer = BytesIO()
    Image.new("RGB", (64, 64), (255, 0, 0)).save(image_buffer, format="WEBP")
    image_path.write_bytes(image_buffer.getvalue())

    image_asset = ArticleImageAsset(
        course_id=course.id,
        article_text_id=article_text.id,
        source_url="https://example.com/cover.webp",
        alt_text="封面图",
        storage_backend="local",
        bucket=None,
        object_key=f"articles/{course.id}/images/cover.webp",
        object_path=str(image_path),
        content_type="image/webp",
        byte_size=image_path.stat().st_size,
        checksum_sha256="a" * 64,
        metadata_json="{}",
        status="imported",
    )
    db_session.add(image_asset)
    db_session.flush()
    article_text.content_markdown = f"# 可下载课程\n\n![封面图](/courses/{course.id}/images/{image_asset.id})"
    db_session.commit()

    original_add_picture = DocxDocument.add_picture

    def guarded_add_picture(self, image_path_or_stream, width=None, height=None):
        if hasattr(image_path_or_stream, "read"):
            position = image_path_or_stream.tell()
            image_path_or_stream.seek(0)
            header = image_path_or_stream.read(8)
            image_path_or_stream.seek(position)
            if header != b"\x89PNG\r\n\x1a\n":
                raise UnrecognizedImageError
        return original_add_picture(self, image_path_or_stream, width=width, height=height)

    monkeypatch.setattr(DocxDocument, "add_picture", guarded_add_picture)
    exported = CourseExportService(db_session).export_content(course, "docx")

    assert exported.data[:2] == b"PK"


def test_course_audio_download_returns_current_audio_attachment(client, db_session, tmp_path):
    course, article_text = create_exportable_course(db_session)
    audio_path = tmp_path / "course.wav"
    audio_path.write_bytes(b"RIFFaudio")
    audio_asset = AudioAsset(
        course_id=course.id,
        article_text_id=article_text.id,
        provider="fake",
        voice_id="fake-cn",
        format="wav",
        object_path=str(audio_path),
        content_type="audio/wav",
        duration_seconds=1,
        character_count=8,
        is_current=True,
    )
    db_session.add(audio_asset)
    db_session.flush()
    course.status = CourseStatus.READY
    course.current_audio_asset_id = audio_asset.id
    db_session.commit()

    response = client.get(f"/courses/{course.id}/audio-download")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("audio/wav")
    assert "attachment" in response.headers["content-disposition"]
    assert ".wav" in response.headers["content-disposition"]
    assert response.content == b"RIFFaudio"
