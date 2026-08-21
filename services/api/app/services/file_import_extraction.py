from __future__ import annotations

import re
import shlex
import subprocess
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from markdownify import markdownify as html_to_markdown

from app.core.config import settings
from app.services.content_normalization import NormalizedContent
from app.services.web_extraction import ExtractionError, extract_article_content, first_markdown_heading


SUPPORTED_EXTENSIONS = {"html", "htm", "pdf", "doc", "docx", "epub", "txt"}
UNSUPPORTED_FORMAT_MESSAGES = {
    "ppt": "PPT import is not supported yet.",
    "pptx": "PPT import is not supported yet.",
    "mobi": "MOBI import is not supported yet.",
    "azw3": "AZW3 import is not supported yet.",
}


class FileImportExtractionError(RuntimeError):
    def __init__(self, message: str, code: str = "file_extraction_failed"):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ExtractedFileContent:
    title: str
    normalized: NormalizedContent
    source_metadata: dict[str, object]
    extraction_metadata: dict[str, object]


def extract_file_content(data: bytes, *, filename: str, content_type: str = "", relative_path: str | None = None) -> ExtractedFileContent:
    extension = file_extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise FileImportExtractionError(
            UNSUPPORTED_FORMAT_MESSAGES.get(extension, "This file type is not supported yet."),
            code="unsupported_file_type",
        )

    locator = relative_path or filename
    if extension in {"html", "htm"}:
        return extract_html_file_content(data, filename=filename, locator=locator, content_type=content_type)
    if extension == "txt":
        return extract_txt_file_content(data, filename=filename, locator=locator, content_type=content_type)
    if extension == "docx":
        return extract_docx_file_content(data, filename=filename, locator=locator, content_type=content_type)
    if extension == "doc":
        return extract_doc_file_content(data, filename=filename, locator=locator, content_type=content_type)
    if extension == "pdf":
        return extract_pdf_file_content(data, filename=filename, locator=locator, content_type=content_type)
    if extension == "epub":
        return extract_epub_file_content(data, filename=filename, locator=locator, content_type=content_type)
    raise FileImportExtractionError("This file type is not supported yet.", code="unsupported_file_type")


def extract_txt_file_content(data: bytes, *, filename: str, locator: str, content_type: str) -> ExtractedFileContent:
    text = decode_text_file(data)
    markdown = normalize_plain_text_to_markdown(text)
    return build_extracted_file_content(
        markdown,
        filename=filename,
        locator=locator,
        content_type=content_type,
        extractor="txt",
        title_override=file_stem(filename),
    )


def extract_html_file_content(data: bytes, *, filename: str, locator: str, content_type: str) -> ExtractedFileContent:
    html = decode_text_file(data)
    pseudo_url = f"https://file-import.local/{Path(locator).name or filename}"
    try:
        extracted = extract_article_content(html, original_url=pseudo_url, final_url=pseudo_url)
        source_metadata = base_source_metadata(filename, locator, content_type)
        extraction_metadata = dict(extracted.extraction_metadata)
        extraction_metadata["extractor"] = f"html_{extracted.extractor}"
        return ExtractedFileContent(
            title=first_markdown_heading(extracted.normalized.content_markdown) or extracted.title or file_stem(filename),
            normalized=extracted.normalized,
            source_metadata=source_metadata,
            extraction_metadata=extraction_metadata,
        )
    except ExtractionError:
        markdown = html_to_markdown(html, heading_style="ATX")
        return build_extracted_file_content(
            markdown,
            filename=filename,
            locator=locator,
            content_type=content_type,
            extractor="html_full_body",
        )


def extract_docx_file_content(data: bytes, *, filename: str, locator: str, content_type: str) -> ExtractedFileContent:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise FileImportExtractionError("Word document could not be parsed.", code="docx_parse_failed") from exc

    blocks: list[str] = []
    title = ""
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if not title:
            title = text
        style_name = getattr(paragraph.style, "name", "")
        if style_name.lower().startswith("heading"):
            blocks.append(f"# {text}")
        else:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return build_extracted_file_content(
        "\n\n".join(blocks),
        filename=filename,
        locator=locator,
        content_type=content_type,
        extractor="docx",
        title_override=title,
    )


def extract_doc_file_content(data: bytes, *, filename: str, locator: str, content_type: str) -> ExtractedFileContent:
    command_template = settings.file_import_doc_converter_command.strip()
    if not command_template:
        raise FileImportExtractionError(
            "Legacy DOC import requires a configured converter.",
            code="legacy_doc_converter_unavailable",
        )

    with TemporaryDirectory() as temp_dir:
        input_path = Path(temp_dir) / "input.doc"
        output_path = Path(temp_dir) / "output.txt"
        input_path.write_bytes(data)
        if "{input}" in command_template or "{output}" in command_template:
            command = shlex.split(command_template.format(input=str(input_path), output=str(output_path)))
        else:
            command = [*shlex.split(command_template), str(input_path)]
        try:
            completed = subprocess.run(command, check=True, capture_output=True, text=True, timeout=60)
        except (OSError, subprocess.SubprocessError) as exc:
            raise FileImportExtractionError(
                "Legacy DOC conversion failed.",
                code="legacy_doc_conversion_failed",
            ) from exc
        text = output_path.read_text(encoding="utf-8", errors="ignore") if output_path.exists() else completed.stdout
    return build_extracted_file_content(
        normalize_plain_text_to_markdown(text),
        filename=filename,
        locator=locator,
        content_type=content_type,
        extractor="doc_converter",
        title_override=first_nonempty_line(text),
    )


def extract_pdf_file_content(data: bytes, *, filename: str, locator: str, content_type: str) -> ExtractedFileContent:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise FileImportExtractionError("PDF could not be parsed.", code="pdf_parse_failed") from exc

    page_texts: list[str] = []
    text_pages = 0
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if re.sub(r"\s+", "", text):
            text_pages += 1
        page_texts.append(text)

    joined_text = "\n\n".join(text for text in page_texts if text.strip())
    non_whitespace_count = len(re.sub(r"\s+", "", joined_text))
    if non_whitespace_count == 0:
        raise FileImportExtractionError(
            "PDF appears to be scanned or image-only. OCR is not supported yet.",
            code="scanned_pdf_without_text_layer",
        )
    if len(page_texts) > 1 and text_pages / max(1, len(page_texts)) < 0.5 and non_whitespace_count < 200:
        raise FileImportExtractionError(
            "PDF text layer is too sparse and appears to be scanned. OCR is not supported yet.",
            code="scanned_pdf_low_text_layer",
        )

    return build_extracted_file_content(
        normalize_plain_text_to_markdown(joined_text),
        filename=filename,
        locator=locator,
        content_type=content_type,
        extractor="pdf_text_layer",
        title_override=first_nonempty_line(joined_text),
        extra_metadata={"page_count": len(page_texts), "text_page_count": text_pages},
    )


def extract_epub_file_content(data: bytes, *, filename: str, locator: str, content_type: str) -> ExtractedFileContent:
    try:
        import ebooklib
        from ebooklib import epub
    except ImportError as exc:
        raise FileImportExtractionError("EPUB parser is not installed.", code="epub_parser_unavailable") from exc

    with TemporaryDirectory() as temp_dir:
        epub_path = Path(temp_dir) / "book.epub"
        epub_path.write_bytes(data)
        try:
            book = epub.read_epub(str(epub_path))
        except Exception as exc:
            raise FileImportExtractionError("EPUB could not be parsed.", code="epub_parse_failed") from exc
    sections: list[str] = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        markdown = html_to_markdown(item.get_content().decode("utf-8", errors="ignore"), heading_style="ATX").strip()
        if markdown:
            sections.append(markdown)

    title = ""
    for value in book.get_metadata("DC", "title"):
        if value and value[0]:
            title = str(value[0]).strip()
            break
    return build_extracted_file_content(
        "\n\n".join(sections),
        filename=filename,
        locator=locator,
        content_type=content_type,
        extractor="epub",
        title_override=title,
    )


def build_extracted_file_content(
    markdown: str,
    *,
    filename: str,
    locator: str,
    content_type: str,
    extractor: str,
    title_override: str = "",
    extra_metadata: dict[str, object] | None = None,
) -> ExtractedFileContent:
    normalized = NormalizedContent.from_markdown(markdown)
    title = title_override.strip() or first_markdown_heading(normalized.content_markdown) or file_stem(filename)
    return ExtractedFileContent(
        title=title,
        normalized=normalized,
        source_metadata=base_source_metadata(filename, locator, content_type),
        extraction_metadata={
            "extractor": extractor,
            "quality": {
                "is_usable": normalized.quality.is_usable,
                "reason": normalized.quality.reason,
                "character_count": normalized.quality.character_count,
            },
            **(extra_metadata or {}),
        },
    )


def base_source_metadata(filename: str, locator: str, content_type: str) -> dict[str, object]:
    return {
        "source_kind": "file",
        "locator": locator,
        "canonical_locator": locator,
        "source_domain": "",
        "original_filename": filename,
        "relative_path": locator,
        "file_extension": file_extension(filename),
        "content_type": content_type,
    }


def decode_text_file(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def normalize_plain_text_to_markdown(text: str) -> str:
    paragraphs = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n\n".join(line for line in paragraphs if line)


def first_nonempty_line(text: str) -> str:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            return stripped
    return ""


def file_extension(filename: str) -> str:
    return Path(filename).suffix.lower().lstrip(".")


def file_stem(filename: str) -> str:
    return Path(filename).stem.strip() or "未命名文件"
