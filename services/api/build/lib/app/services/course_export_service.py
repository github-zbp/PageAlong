from __future__ import annotations

import html
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

import httpx
from docx import Document
from docx.shared import Inches
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Image, Paragraph, Preformatted, SimpleDocTemplate, Spacer
from reportlab.lib.utils import ImageReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.course import ArticleImageAsset, ArticleText, Course

IMAGE_MARKDOWN_PATTERN = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)$")
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+)$")
LIST_ITEM_PATTERN = re.compile(r"^[-*+]\s+(.+)$")


@dataclass(frozen=True)
class CourseExport:
    data: bytes
    media_type: str
    filename: str


class CourseExportService:
    def __init__(self, db: Session):
        self.db = db

    def export_content(self, course: Course, export_format: str) -> CourseExport:
        normalized_format = export_format.strip().lower()
        article_text = self._latest_article_text(course)
        markdown = self._markdown_for_course(course, article_text)
        filename_base = safe_filename(course.title or "course")

        if normalized_format in {"markdown", "md"}:
            return CourseExport(
                data=markdown.encode("utf-8"),
                media_type="text/markdown; charset=utf-8",
                filename=f"{filename_base}.md",
            )
        if normalized_format == "docx":
            return CourseExport(
                data=self._build_docx(course, markdown),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                filename=f"{filename_base}.docx",
            )
        if normalized_format == "pdf":
            return CourseExport(
                data=self._build_pdf(course, markdown),
                media_type="application/pdf",
                filename=f"{filename_base}.pdf",
            )
        raise ValueError("Unsupported export format")

    def _latest_article_text(self, course: Course) -> ArticleText | None:
        return self.db.scalar(
            select(ArticleText)
            .where(ArticleText.course_id == course.id)
            .order_by(ArticleText.version.desc(), ArticleText.created_at.desc())
        )

    def _markdown_for_course(self, course: Course, article_text: ArticleText | None) -> str:
        if article_text is not None and article_text.content_markdown.strip():
            return article_text.content_markdown.strip() + "\n"
        if article_text is not None and article_text.text.strip():
            return f"# {course.title}\n\n{article_text.text.strip()}\n"
        sentence_text = "\n\n".join(sentence.text for sentence in sorted(course.sentences, key=lambda item: item.index))
        return f"# {course.title}\n\n{sentence_text.strip()}\n"

    def _build_docx(self, course: Course, markdown: str) -> bytes:
        document = Document()
        document.core_properties.title = course.title
        for block in parse_markdown_blocks(markdown):
            kind = block["kind"]
            text = block["text"]
            if kind == "heading":
                document.add_heading(text, level=min(int(block["level"]), 4))
            elif kind == "list":
                document.add_paragraph(text, style="List Bullet")
            elif kind == "code":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(text)
                run.font.name = "Courier New"
            elif kind == "image":
                image_data = self._image_data(course, block["src"])
                if image_data is not None:
                    document.add_picture(BytesIO(image_data), width=Inches(5.8))
                if text:
                    document.add_paragraph(text)
            else:
                document.add_paragraph(text)
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def _build_pdf(self, course: Course, markdown: str) -> bytes:
        output = BytesIO()
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        styles = getSampleStyleSheet()
        body_style = ParagraphStyle(
            "PageAlongBody",
            parent=styles["BodyText"],
            fontName="STSong-Light",
            fontSize=11,
            leading=17,
            alignment=TA_LEFT,
            spaceAfter=8,
        )
        heading_style = ParagraphStyle(
            "PageAlongHeading",
            parent=body_style,
            fontSize=18,
            leading=24,
            spaceBefore=8,
            spaceAfter=12,
        )
        code_style = ParagraphStyle(
            "PageAlongCode",
            parent=body_style,
            fontName="Courier",
            fontSize=9,
            leading=12,
        )
        story: list[object] = []
        for block in parse_markdown_blocks(markdown):
            kind = block["kind"]
            text = html.escape(block["text"])
            if kind == "heading":
                story.append(Paragraph(text, heading_style))
            elif kind == "code":
                story.append(Preformatted(block["text"], code_style))
            elif kind == "image":
                image_flowable = self._pdf_image(course, block["src"])
                if image_flowable is not None:
                    story.append(image_flowable)
                    story.append(Spacer(1, 8))
                if text:
                    story.append(Paragraph(text, body_style))
            else:
                story.append(Paragraph(text, body_style))
        if not story:
            story.append(Paragraph(html.escape(course.title), heading_style))
        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=0.7 * inch,
            leftMargin=0.7 * inch,
            topMargin=0.7 * inch,
            bottomMargin=0.7 * inch,
            title=course.title,
        )
        document.build(story)
        return output.getvalue()

    def _pdf_image(self, course: Course, src: str) -> Image | None:
        data = self._image_data(course, src)
        if data is None:
            return None
        try:
            reader = ImageReader(BytesIO(data))
            width, height = reader.getSize()
        except Exception:
            return None
        max_width = 6.0 * inch
        scale = min(1.0, max_width / max(1, width))
        return Image(BytesIO(data), width=width * scale, height=height * scale)

    def _image_data(self, course: Course, src: str) -> bytes | None:
        asset = self._image_asset_for_source(course, src)
        if asset is not None:
            return self._image_asset_data(asset)
        if src.lower().startswith(("http://", "https://")):
            try:
                response = httpx.get(src, timeout=8)
                response.raise_for_status()
                return response.content
            except Exception:
                return None
        return None

    def _image_asset_for_source(self, course: Course, src: str) -> ArticleImageAsset | None:
        parsed = urlparse(src)
        path = parsed.path if parsed.scheme else src
        marker = f"/courses/{course.id}/images/"
        if marker not in path:
            return None
        image_asset_id = path.rsplit("/", 1)[-1]
        asset = self.db.get(ArticleImageAsset, image_asset_id)
        if asset is None or asset.course_id != course.id or asset.status != "imported":
            return None
        return asset

    def _image_asset_data(self, asset: ArticleImageAsset) -> bytes | None:
        if asset.storage_backend == "local":
            path = Path(asset.object_path)
            if path.exists():
                return path.read_bytes()
            return None
        if asset.object_path.lower().startswith(("http://", "https://")):
            try:
                response = httpx.get(asset.object_path, timeout=8)
                response.raise_for_status()
                return response.content
            except Exception:
                return None
        if asset.storage_backend in {"s3", "minio", "r2"}:
            try:
                import boto3

                response = boto3.client(
                    "s3",
                    endpoint_url=settings.s3_endpoint_url,
                    aws_access_key_id=settings.s3_access_key_id,
                    aws_secret_access_key=settings.s3_secret_access_key,
                ).get_object(Bucket=asset.bucket or settings.s3_bucket, Key=asset.object_key)
                return response["Body"].read()
            except Exception:
                return None
        return None


def parse_markdown_blocks(markdown: str) -> list[dict[str, str]]:
    blocks: list[dict[str, str]] = []
    paragraph: list[str] = []
    code_lines: list[str] = []
    in_code = False

    def flush_paragraph() -> None:
        nonlocal paragraph
        if paragraph:
            blocks.append({"kind": "paragraph", "text": " ".join(paragraph).strip()})
            paragraph = []

    for line in markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                blocks.append({"kind": "code", "text": "\n".join(code_lines)})
                code_lines = []
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not stripped:
            flush_paragraph()
            continue
        heading = HEADING_PATTERN.match(stripped)
        if heading:
            flush_paragraph()
            blocks.append({"kind": "heading", "level": str(len(heading.group(1))), "text": heading.group(2)})
            continue
        image = IMAGE_MARKDOWN_PATTERN.match(stripped)
        if image:
            flush_paragraph()
            blocks.append({"kind": "image", "text": image.group(1), "src": image.group(2)})
            continue
        list_item = LIST_ITEM_PATTERN.match(stripped)
        if list_item:
            flush_paragraph()
            blocks.append({"kind": "list", "text": list_item.group(1)})
            continue
        paragraph.append(stripped)
    flush_paragraph()
    if in_code and code_lines:
        blocks.append({"kind": "code", "text": "\n".join(code_lines)})
    return blocks


def safe_filename(value: str) -> str:
    normalized = re.sub(r"[^\w\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff.-]+", "_", value.strip(), flags=re.UNICODE)
    normalized = normalized.strip("._")
    return (normalized or "course")[:80]
