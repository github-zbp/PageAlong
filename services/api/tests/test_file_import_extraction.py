from io import BytesIO

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app.services.file_import_extraction import FileImportExtractionError, extract_file_content


def test_extract_txt_decodes_gb18030_and_normalizes_text():
    data = "第一段正文。\n\n第二段正文。".encode("gb18030")

    result = extract_file_content(data, filename="notes.txt", content_type="text/plain")

    assert result.title == "notes"
    assert "第一段正文" in result.normalized.content_markdown
    assert "第二段正文" in result.normalized.tts_text
    assert result.source_metadata["file_extension"] == "txt"


def test_extract_html_uses_existing_article_cleanup():
    html = """
    <html><head><title>网页文件</title></head><body>
      <nav>首页 登录</nav>
      <article>
        <h1>正文标题</h1>
        <p>第一段正文内容足够长，用来验证 HTML 文件会走统一清洗流程。</p>
        <p>第二段正文继续补充内容，避免质量检测误判为正文过短。</p>
        <p>第三段正文继续补充足够字符数量，让提取结果稳定可用。</p>
      </article>
    </body></html>
    """.encode()

    result = extract_file_content(html, filename="article.html", content_type="text/html")

    assert result.title == "正文标题"
    assert "正文标题" in result.normalized.content_markdown
    assert "首页 登录" not in result.normalized.content_markdown
    assert result.extraction_metadata["extractor"].startswith("html_")


def test_extract_docx_reads_paragraphs_and_table_cells():
    document = Document()
    document.add_heading("文档标题", level=1)
    document.add_paragraph("第一段正文。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "表格左侧。"
    table.cell(0, 1).text = "表格右侧。"
    buffer = BytesIO()
    document.save(buffer)

    result = extract_file_content(
        buffer.getvalue(),
        filename="lesson.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result.title == "文档标题"
    assert "第一段正文" in result.normalized.tts_text
    assert "表格左侧" in result.normalized.tts_text
    assert "表格右侧" in result.normalized.tts_text


def test_extract_pdf_rejects_scanned_or_image_only_pdf():
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.rect(100, 500, 200, 100, stroke=1, fill=0)
    pdf.showPage()
    pdf.save()

    with pytest.raises(FileImportExtractionError) as exc_info:
        extract_file_content(buffer.getvalue(), filename="scan.pdf", content_type="application/pdf")

    assert exc_info.value.code == "scanned_pdf_without_text_layer"


def test_extract_unsupported_formats_fail_explicitly():
    with pytest.raises(FileImportExtractionError) as exc_info:
        extract_file_content(b"slides", filename="slides.pptx", content_type="application/vnd.ms-powerpoint")

    assert exc_info.value.code == "unsupported_file_type"
